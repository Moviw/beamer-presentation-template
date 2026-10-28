#!/usr/bin/env python3
"""从最终 PDF 导出 .pptx：每页渲染成高清图，铺满一张 16:9 幻灯片（不重新排版，和 PDF 视觉一致）。
可选 --key：在 macOS 上用 Keynote 把 .pptx 另存为 .key（Linux 上没有 Keynote，只出 .pptx）。
用法：python3 pdf2pptx.py presentation.pdf [--dpi 200] [--out presentation.pptx] [--key]
依赖：poppler（pdftoppm）、python-pptx（pip install python-pptx）"""
import argparse, glob, os, platform, subprocess, sys, tempfile

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:
    sys.exit("缺少 python-pptx：pip install python-pptx")

ap = argparse.ArgumentParser()
ap.add_argument("pdf")
ap.add_argument("--dpi", type=int, default=200)
ap.add_argument("--out")
ap.add_argument("--key", action="store_true", help="macOS：再用 Keynote 另存为 .key")
a = ap.parse_args()
out = os.path.abspath(a.out or os.path.splitext(a.pdf)[0] + ".pptx")

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)      # 16:9（13.333 × 7.5 英寸）
with tempfile.TemporaryDirectory() as d:
    subprocess.run(["pdftoppm", "-png", "-r", str(a.dpi), a.pdf, f"{d}/p"], check=True)
    pages = sorted(glob.glob(f"{d}/p-*.png"))
    for png in pages:
        s = prs.slides.add_slide(prs.slide_layouts[6])                 # 空白版式
        s.shapes.add_picture(png, 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(out)
print(f"{len(pages)} 页 → {out}")

if a.key:
    if platform.system() != "Darwin":
        sys.exit("--key 需要 macOS 上的 Keynote；这里只生成了 .pptx")
    key = out[:-5] + ".key"
    subprocess.run(["osascript", "-e", f'''tell application "Keynote"
        set d to open POSIX file "{out}"
        delay 3
        save d in POSIX file "{key}"
        close d
    end tell'''], check=True)
    print(f"→ {key}")
